"""Generate an accountant-ready Excel workbook with executive summary, discrepancy details, and anomalies."""
from __future__ import annotations
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
REGULAR_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", size=10, bold=True)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=False)
LEFT_ALIGN = Alignment(horizontal="left", vertical="center")
RIGHT_ALIGN = Alignment(horizontal="right", vertical="center")

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9")
)

GREEN_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
RED_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

def _style_worksheet(ws, is_summary=False):
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Style Header Row
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN

    # Style Data Rows
    for row_idx in range(2, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            val = cell.value

            # Number formatting for currency floats
            if isinstance(val, (int, float)) and not isinstance(val, bool):
                cell.number_format = "#,##0.00"
                cell.alignment = RIGHT_ALIGN
            elif isinstance(val, str) and val in ("✓ Сошлось", "OK", "MATCHED"):
                cell.fill = GREEN_FILL
                cell.font = BOLD_FONT
                cell.alignment = CENTER_ALIGN
            elif isinstance(val, str) and val in ("Требует внимания", "Высокий", "Критический"):
                cell.fill = RED_FILL
                cell.font = BOLD_FONT
                cell.alignment = CENTER_ALIGN
            elif isinstance(val, str) and len(val) <= 10:
                cell.alignment = CENTER_ALIGN
            else:
                cell.alignment = LEFT_ALIGN

    # Auto-adjust column width
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(46, max(13, max_len + 3))

def run(reconciliation: pd.DataFrame, discrepancies: pd.DataFrame, anomalies: pd.DataFrame, output_dir: str | Path):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    report_path = output_dir / "reconciliation_report.xlsx"

    with pd.ExcelWriter(report_path, engine="openpyxl") as writer:
        reconciliation.to_excel(writer, sheet_name="Сводка по субагентам", index=False)
        discrepancies.to_excel(writer, sheet_name="Детализация расхождений", index=False)
        anomalies.to_excel(writer, sheet_name="Аномалии и антифрод", index=False)

        book = writer.book
        for name in book.sheetnames:
            ws = book[name]
            _style_worksheet(ws, is_summary=(name == "Сводка по субагентам"))

    return report_path
