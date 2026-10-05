"""Excel report for the accountant and the jury. Totals and the bridge identity are live formulas."""
from __future__ import annotations

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .reconcile import COMPONENTS

FONT = "Arial"
HDR_FILL = PatternFill("solid", start_color="1F3864")
SUB_FILL = PatternFill("solid", start_color="D9E2F3")
NUM = "#,##0.00;[Red]-#,##0.00;-"
INT = "#,##0"

STATUS_RU = {
    "AMOUNT_RESIDUAL": "Сумма не сходится", "DOUBLE_DEBIT": "Двойное списание в ETM",
    "NOT_IN_1C": "Есть в ETM, нет в 1С", "ONEC_ONLY_REGISTRY_SALE": "Есть в реестре и 1С, нет в ETM",
    "ONEC_ONLY": "Есть только в 1С", "ETM_IN_MISSING_ACT_MONTH": "ETM, акт 1С за месяц не выгружен",
    "VOID_UNBALANCED": "Войд не закрывает выкуп", "REGISTRY_AMOUNT_DIFF": "Реестр отличается от ETM и 1С", "REGISTRY_DUPLICATE_ROW": "Строка реестра задвоена",
    "REGISTRY_TICKET_IN_AMOUNT": "В реестре билет вместо суммы", "REGISTRY_ONLY_SALE": "Продажа только в реестре",
    "REGISTRY_ONLY_OTHER": "Возврат/войд только в реестре",
}
SRC_RU = {"1C": "1С", "BOT": "Бот ETM", "AGENT": "Реестр (сотрудник)", "EXPORT": "Выгрузка 1С", "REVIEW": "Нужна проверка", "OK": "OK"}
WHO = {"1C": "Бухгалтер 1С", "BOT": "Разработчик/поддержка бота ETM", "AGENT": "Сотрудник, ведущий реестр",
       "EXPORT": "Бухгалтер: выгрузить акт", "REVIEW": "Бухгалтер + менеджер субагента", "OK": "-"}


def _clean(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime().date() if (v.hour == 0 and v.minute == 0) else v.to_pydatetime()
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        return float(v)
    if isinstance(v, np.bool_):
        return bool(v)
    return v


def _hdr(cell):
    cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    cell.fill = HDR_FILL
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def write_table(ws, df, start_row=1, widths=None, money_cols=(), int_cols=(), wrap_cols=(), freeze_col=2):
    for j, c in enumerate(df.columns, 1):
        _hdr(ws.cell(start_row, j, c))
    ws.row_dimensions[start_row].height = 42
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, 1):
            cell = ws.cell(i, j, _clean(v))
            cell.font = Font(name=FONT, size=10)
            col = df.columns[j - 1]
            if col in money_cols:
                cell.number_format = NUM
            elif col in int_cols:
                cell.number_format = INT
            elif hasattr(cell.value, "year"):
                cell.number_format = "DD.MM.YYYY"
            if col in wrap_cols:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for j, c in enumerate(df.columns, 1):
        w = (widths or {}).get(c)
        if w is None:
            sample = df[c].head(200).map(lambda x: len(str(x)) if x is not None and x == x else 0)
            w = min(max(11, int(sample.max() if len(sample) else 10) + 2, len(str(c)) // 2 + 4), 46)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(start_row + 1, freeze_col) if freeze_col else None
    if len(df):
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(df.columns))}{start_row + len(df)}"
    return start_row + len(df)


def _title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=FONT, italic=True, size=10, color="555555")


def build_cases(R):
    g = R.groups
    cases = []
    c1 = g[(g.status == "AMOUNT_RESIDUAL") & (g.error_source == "1C") & (g.confidence == "A") & (g.reg_rows > 0) & (g.ticket_count == 1)]
    c1 = c1.assign(a=c1.residual.abs()).sort_values("a", ascending=False).head(1)
    c2 = g[g.status == "DOUBLE_DEBIT"]
    c2 = c2.assign(a=c2.dup_extra_amount.abs()).sort_values("a", ascending=False).head(1)
    pay = R.payments[R.payments.status == "REPEATED_CREDIT"]
    pay = pay.assign(a=pay.amount_kgs.abs()).sort_values("a", ascending=False).head(1)
    for r in c1.itertuples():
        cases.append(("Кейс 1. Неверная сумма в 1С (подтверждена реестром)", r.agent,
                      f"Билет {r.tickets}: ETM списал {-r.etm_sum:,.2f}, реестр {r.reg_sum:,.2f}, а в 1С реализация на {r.act_sum:,.2f}.",
                      f"Реестр и ETM совпадают, 1С одна отличается на {abs(r.residual):,.2f} ({r.act_docs}).",
                      f"Бухгалтер 1С; запись в реестре вёл {r.employees or '—'}.",
                      f"Исправить реализацию в 1С на {-r.residual:+,.2f}; баланс ETM не трогать."))
    for r in c2.itertuples():
        cases.append(("Кейс 2. Бот списал билет дважды", r.agent,
                      f"Билет {r.tickets}: в ETM два одинаковых списания, лишнее {-r.dup_extra_amount:,.2f}.",
                      f"ETM-операции {r.etm_txn_ids}; без дубля сумма ETM = 1С ({r.act_sum:,.2f}).",
                      "Разработчик/поддержка бота ETM (защита от повторной проводки); менеджер субагента — возврат списания.",
                      f"Вернуть на баланс ETM {-r.dup_extra_amount:,.2f}."))
    for r in pay.itertuples():
        cases.append(("Кейс 3. Оплата зачислена в ETM повторно", r.agent,
                      f"Оплата {r.amount_kgs:,.2f} от {r.date:%d.%m.%Y}: в ETM две одинаковые проводки, в 1С поступление одно.",
                      r.explanation, "Разработчик/поддержка бота ETM; казначей (подтвердить, что деньги пришли один раз).",
                      f"Списать с баланса ETM {r.amount_kgs:,.2f}."))
    return cases


def build_workbook(R, M, emp, summary, path):
    wb = Workbook()
    act = R.actions
    unexplained = R.bridge.set_index("agent")["unexplained"]

    # ------------------------------------------------------------ По субагентам
    ws = wb.active
    ws.title = "По субагентам"
    _title(ws, "Отчёт по каждому субагенту: что двигать, что исправлять, что выяснить",
           "Разрыв = сальдо 1С + баланс ETM (идеал 0). Корректировки закрывают найденные расхождения; "
           "то, что данные объяснить не могут, вынесено в «Выяснить», а не растворено в остатке.")
    heads = ["Субагент", "Разрыв на начало периода", "Разрыв на конец (сальдо 1С + баланс ETM)",
             "Сдвинуть баланс ETM на, KGS", "Исправить в 1С на, KGS", "Сначала выяснить, KGS", "Округления групп, KGS"]
    d = pd.DataFrame({heads[0]: act["agent"], heads[1]: act["opening_gap"], heads[2]: act["gap_close"],
                      heads[3]: act["etm_adjust_kgs"], heads[4]: act["onec_adjust_kgs"],
                      heads[5]: act["to_clarify_kgs"], heads[6]: act["rounding_kgs"]})
    d["Рекомендация"] = act["recommended_action"]
    skip = {"n_OK", "n_VOID_CLOSED", "n_REGISTRY_DIRECT_CLIENT"}
    d["Проблемных групп и платежей"] = act[[c for c in act.columns if c.startswith("n_") and c not in skip]].sum(axis=1)
    d["Месяцев без акта 1С"] = act["missing_1c_months"]
    r0 = 4
    last = write_table(ws, d, r0, widths={"Рекомендация": 90, "Субагент": 30}, money_cols=heads[1:],
                       wrap_cols=["Рекомендация"], int_cols=["Проблемных групп и платежей", "Месяцев без акта 1С"])
    c_after, c_chk = len(d.columns) + 1, len(d.columns) + 2
    ws.cell(r0, c_after, "Остаток после правок, KGS (=C+D+E)")
    ws.cell(r0, c_chk, "Проверка: остаток − выяснить − округление − необъяснённое (≈0)")
    for c in (c_after, c_chk):
        _hdr(ws.cell(r0, c))
        ws.column_dimensions[get_column_letter(c)].width = 24
    A, K = get_column_letter(c_after), None
    for i in range(r0 + 1, last + 1):
        un = unexplained[d.iloc[i - r0 - 1, 0]]
        ws.cell(i, c_after, f"=C{i}+D{i}+E{i}").number_format = NUM
        ws.cell(i, c_chk, f"=ROUND({A}{i}-F{i}-G{i}-({un:.9f}),2)").number_format = NUM
        for c in (c_after, c_chk):
            ws.cell(i, c).font = Font(name=FONT, size=10)
    tot = last + 1
    ws.cell(tot, 1, "ИТОГО").font = Font(name=FONT, bold=True)
    ws.cell(tot, 1).fill = SUB_FILL
    for col in list(range(2, 8)) + [c_after]:
        L = get_column_letter(col)
        ws.cell(tot, col, f"=SUM({L}{r0 + 1}:{L}{last})").number_format = NUM
        ws.cell(tot, col).font = Font(name=FONT, bold=True)
    ws.freeze_panes = ws.cell(r0 + 1, 2)
    SUB = dict(first=r0 + 1, last=last, tot=tot, after=A)

    # ------------------------------------------------------------------ Мост
    ws = wb.create_sheet("Мост")
    _title(ws, "Баланс-мост по субагентам (без заплаток): разрыв на конец = начальный разрыв + сумма найденных причин",
           "Каждая строка 1С и ETM попала ровно в одну колонку. «Не объяснено» — формула. Колонки «акт не выгружен» и «скачок через пропуск» — "
           "честно нераспределённые суммы, а не ошибки.")
    comp_ru = {
        "opening_gap": "Разрыв на начало", "onec_continuity_breaks": "1С: разрыв между актами",
        "onec_jumps_over_missing_acts": "1С: скачок через пропущенный акт", "onec_act_arithmetic": "1С: арифметика акта",
        "rounding_matched_groups": "Округления сошедшихся групп", "amount_residual_1c": "Сумма: ошибка 1С",
        "amount_residual_bot": "Сумма: ошибка бота", "amount_residual_review": "Сумма: нужна проверка",
        "double_debit": "Двойное списание", "onec_only_registry_sale": "Бот не перенёс продажу",
        "onec_only_no_registry": "Только 1С, без реестра", "etm_not_in_1c": "Есть в ETM, нет в 1С",
        "void_unbalanced": "Войд не закрыт", "etm_in_missing_act_month": "ETM: акт не выгружен",
        "payment_not_in_1c": "Оплата не проведена в 1С", "payment_not_in_etm": "Оплата не зачислена в ETM",
        "repeated_credit": "Повторная оплата", "payment_wrong_agent": "Оплата не тому субагенту",
        "payment_in_missing_act_month": "Оплата: акт не выгружен",
    }
    bb = R.bridge[["agent", "gap_close"] + COMPONENTS].rename(columns={"agent": "Субагент", "gap_close": "Разрыв на конец", **comp_ru})
    last = write_table(ws, bb, 3, money_cols=list(bb.columns[1:]), widths={"Субагент": 30})
    nc = len(bb.columns)
    sumc, unc = nc + 1, nc + 2
    ws.cell(3, sumc, "Сумма причин")
    ws.cell(3, unc, "Не объяснено (формула)")
    for c in (sumc, unc):
        _hdr(ws.cell(3, c))
        ws.column_dimensions[get_column_letter(c)].width = 20
    c1, c2 = get_column_letter(3), get_column_letter(nc)
    UNC = get_column_letter(unc)
    for i in range(4, last + 1):
        ws.cell(i, sumc, f"=SUM({c1}{i}:{c2}{i})").number_format = NUM
        ws.cell(i, unc, f"=ROUND(B{i}-{get_column_letter(sumc)}{i},4)").number_format = "0.0000"
    tot = last + 1
    ws.cell(tot, 1, "ИТОГО").font = Font(name=FONT, bold=True)
    for col in range(2, sumc + 1):
        L = get_column_letter(col)
        ws.cell(tot, col, f"=SUM({L}4:{L}{last})").number_format = NUM
        ws.cell(tot, col).font = Font(name=FONT, bold=True)
    ws.cell(tot, unc, f"=SUM({UNC}4:{UNC}{last})").number_format = "0.0000"
    ws.freeze_panes = "C4"
    BR = dict(last=last, unc=UNC)

    # ---------------------------------------------------------------- Итог
    s = wb.create_sheet("Итог", 0)
    _title(s, "Сверка ETM ↔ 1С ↔ реестр: итог", "Числа в первом блоке — формулы по листам отчёта.")
    rows = [
        ("Субагентов в сверке", f"=COUNTA('По субагентам'!A{SUB['first']}:A{SUB['last']})", INT),
        ("Суммарный разрыв (сальдо 1С + баланс ETM) на конец периода, KGS", f"='По субагентам'!C{SUB['tot']}", NUM),
        ("Рекомендуется сдвинуть баланс ETM (нетто), KGS", f"='По субагентам'!D{SUB['tot']}", NUM),
        ("Рекомендуется исправить в 1С (нетто), KGS", f"='По субагентам'!E{SUB['tot']}", NUM),
        ("Нужно выяснить (нетто), KGS", f"='По субагентам'!F{SUB['tot']}", NUM),
        ("Остаток разрыва после правок (=выяснить + округления), KGS", f"='По субагентам'!{SUB['after']}{SUB['tot']}", NUM),
        ("Не объяснено мостом: сумма по 100 субагентам, KGS", f"=SUM('Мост'!{BR['unc']}4:{BR['unc']}{BR['last']})", "0.0000"),
    ]
    for i, (k, f, fmt) in enumerate(rows, 4):
        s.cell(i, 1, k).font = Font(name=FONT, size=11)
        c = s.cell(i, 2, f)
        c.number_format = fmt
        c.font = Font(name=FONT, size=11, bold=True)
    i = 4 + len(rows) + 1
    s.cell(i, 1, "Что нашли (факты по данным)").font = Font(name=FONT, bold=True, size=12)
    i += 1
    for k, v in summary["headline"]:
        s.cell(i, 1, k).font = Font(name=FONT, size=10)
        c = s.cell(i, 2, v)
        c.font = Font(name=FONT, size=10, bold=True)
        c.alignment = Alignment(horizontal="right")
        if isinstance(v, (int, np.integer)):
            c.number_format = INT
        elif isinstance(v, float):
            c.number_format = NUM
        i += 1
    i += 1
    s.cell(i, 1, "Честные оговорки").font = Font(name=FONT, bold=True, size=12)
    for t in summary["caveats"]:
        i += 1
        s.cell(i, 1, "• " + t).font = Font(name=FONT, size=10)
        s.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
        s.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
        s.row_dimensions[i].height = 44
    s.column_dimensions["A"].width = 110
    s.column_dimensions["B"].width = 24

    # ------------------------------------------------------------ Мост по месяцам
    ws = wb.create_sheet("Мост по месяцам")
    _title(ws, "Помесячная проверка тождества (без подгонки)",
           "Изменение разрыва за месяц обязано равняться: арифметика акта + оборот 1С + оборот ETM. «Остаток тождества» — формула.")
    mo = R.monthly[["agent", "month", "saldo_start_1c", "saldo_end_1c", "movement_1c", "etm_balance_start", "etm_balance_end",
                    "movement_etm", "gap_start", "gap_end", "act_internal_delta", "continuity_break", "jump_in_1c_saldo"]].copy()
    mo.columns = ["Субагент", "Месяц", "Сальдо 1С начало", "Сальдо 1С конец", "Оборот 1С", "Баланс ETM начало", "Баланс ETM конец",
                  "Оборот ETM", "Разрыв начало", "Разрыв конец", "Арифметика акта", "Разрыв сальдо между актами?", "Скачок сальдо 1С на входе"]
    last = write_table(ws, mo, 3, money_cols=list(mo.columns[2:11]) + [mo.columns[12]], widths={"Субагент": 30})
    c = len(mo.columns) + 1
    ws.cell(3, c, "Остаток тождества (≈0)")
    _hdr(ws.cell(3, c))
    ws.column_dimensions[get_column_letter(c)].width = 22
    for i in range(4, last + 1):
        ws.cell(i, c, f"=ROUND((J{i}-I{i})-(K{i}+E{i}+H{i}),4)").number_format = "0.0000"

    # ---------------------------------------------------------------- Расхождения
    ws = wb.create_sheet("Расхождения")
    g = R.groups
    bad = g[~g.status.isin(["OK", "VOID_CLOSED", "REGISTRY_DIRECT_CLIENT"])].copy()
    bad["abs"] = bad["residual"].abs()
    bad = bad.sort_values(["confidence", "abs"], ascending=[True, False])
    ex = pd.DataFrame({
        "ID": bad["group_id"], "Субагент": bad["agent"], "Билеты": bad["tickets"],
        "Тип": bad["status"].map(STATUS_RU).fillna(bad["status"]), "Источник ошибки": bad["error_source"].map(SRC_RU),
        "Уверенность": bad["confidence"], "Сумма 1С": bad["act_sum"], "Сумма ETM": bad["etm_sum"], "Сумма реестра": bad["reg_sum"],
        "Остаток 1С+ETM": bad["residual"], "Дата ETM": bad["etm_date_min"], "Дата 1С": bad["act_date_min"],
        "Паттерн суммы": bad["amount_pattern"], "Кто проверяет": bad["error_source"].map(WHO), "Сотрудник реестра": bad["employees"],
        "Что произошло": bad["explanation"], "ETM txn": bad["etm_txn_ids"],
    })
    _title(ws, f"Расхождения по группам билетов ({len(ex):,} строк — найденные случаи с источником и уверенностью, а не «ошибки»)",
           "Уверенность: A — подтверждено вторым источником; B — правдоподобная гипотеза; C — данных недостаточно.")
    write_table(ws, ex, 3, money_cols=["Сумма 1С", "Сумма ETM", "Сумма реестра", "Остаток 1С+ETM"],
                widths={"Что произошло": 80, "Билеты": 24, "Субагент": 28, "Кто проверяет": 30}, wrap_cols=["Что произошло"])

    # ----------------------------------------------------------------- Платежи
    ws = wb.create_sheet("Платежи")
    p = R.payments
    pe = pd.DataFrame({"Субагент": p["agent"], "Дата": p["date"], "Где найдено": p["where"], "Документ/ID": p["reference"].astype(str),
                       "Сумма": p["amount_kgs"], "Тип": p["status"], "Источник": p["error_source"].map(SRC_RU),
                       "Уверенность": p["confidence"], "Что произошло": p["explanation"], "Влияние на разрыв": p["gap_effect"]})
    _title(ws, "Платежи: оплаты ETM без поступления в 1С и поступления 1С без оплаты в ETM")
    write_table(ws, pe.sort_values("Влияние на разрыв", key=abs, ascending=False), 3, money_cols=["Сумма", "Влияние на разрыв"],
                widths={"Что произошло": 85, "Субагент": 28}, wrap_cols=["Что произошло"])

    # ----------------------------------------------------------------- Аномалии
    ws = wb.create_sheet("Аномалии")
    dd = g[g.status == "DOUBLE_DEBIT"]
    parts = [pd.DataFrame({"Тип": "Двойное списание", "Субагент": dd["agent"], "Билеты/документ": dd["tickets"], "Дата": dd["etm_date_min"],
                           "Сумма, KGS": dd["dup_extra_amount"], "Описание": dd["explanation"]})]
    for st, lab in [("REPEATED_CREDIT", "Повторная оплата"), ("PAYMENT_WRONG_AGENT", "Оплата не тому субагенту")]:
        q = p[p.status == st]
        parts.append(pd.DataFrame({"Тип": lab, "Субагент": q["agent"], "Билеты/документ": q["reference"].astype(str), "Дата": q["date"],
                                   "Сумма, KGS": q["amount_kgs"], "Описание": q["explanation"]}))
    an = pd.concat(parts).sort_values("Сумма, KGS", key=abs, ascending=False)
    _title(ws, f"Повторные операции и аномалии ({len(an)} шт.)")
    write_table(ws, an, 3, money_cols=["Сумма, KGS"], widths={"Описание": 90, "Субагент": 28}, wrap_cols=["Описание"])

    # ---------------------------------------------------------------- 1С
    ws = wb.create_sheet("1С разрывы между актами")
    at = R.act_table
    br = at[at.continuity_break | at.jump_over_missing].copy()
    cb = pd.DataFrame({"Субагент": br["agent"], "Период": br["period_start"].dt.strftime("%Y-%m"),
                       "Конечное сальдо пред. акта": br["prev_end"], "Начальное сальдо акта": br["saldo_start"], "Скачок": br["jump"],
                       "Тип": np.where(br["continuity_break"], "Разрыв (месяцы подряд) — ошибка 1С", "Через пропущенный месяц — акт не выгружен")})
    _title(ws, "1С: непрерывность сальдо между актами",
           "Повторяющиеся суммы у разных субагентов (107 557,10; 283 308,99 …) — вопрос бухгалтеру: пакетная корректировка?")
    write_table(ws, cb.sort_values("Скачок", key=abs, ascending=False), 3,
                money_cols=["Конечное сальдо пред. акта", "Начальное сальдо акта", "Скачок"], widths={"Субагент": 30, "Тип": 50})

    ws = wb.create_sheet("1С нет актов")
    _title(ws, "Субагент-месяцы без акта 1С (из-за них часть сверки невозможна)")
    write_table(ws, R.missing.rename(columns={"agent": "Субагент", "month": "Месяц"}), 3, widths={"Субагент": 34})

    ws = wb.create_sheet("1С черновик и переиздан")
    _title(ws, "Версии актов: используется последняя (переиздан); ниже — чем версии различаются",
           "Каждая версия по отдельности арифметически сходится; склеивать их нельзя — это и давало ложные «нарушения целостности».")
    v = R.versions.rename(columns={
        "folder": "Папка 1С", "period": "Период", "saldo_end_draft": "Сальдо конец (черновик)", "saldo_end_reissued": "Сальдо конец (переиздан)",
        "delta_saldo_end": "Разница сальдо", "rows_draft": "Строк (черновик)", "rows_reissued": "Строк (переиздан)",
        "rows_only_in_reissued": "Только в переизданном", "rows_only_in_draft": "Только в черновике",
        "draft_saldo_check": "Проверка арифметики (черновик)", "reissued_saldo_check": "Проверка арифметики (переиздан)"})
    write_table(ws, v, 3, money_cols=[c for c in v.columns if "альдо" in c or "Проверка" in c or "Разница" in c], widths={"Папка 1С": 34})

    # -------------------------------------------------------------- Сотрудники/модель
    ws = wb.create_sheet("Сотрудники реестра")
    _title(ws, "Ошибки записи в реестре по сотрудникам (продажи по субагентам)",
           "Различие значимо, только если p-value < 0.05. Если нет — честный вывод: концентрации ошибок по сотруднику в данных нет.")
    e2 = emp.rename(columns={"employee": "Сотрудник", "errors": "Ошибок в реестре", "rows": "Продаж", "error_rate": "Доля ошибок",
                             "ci95_low": "95% ДИ от", "ci95_high": "95% ДИ до", "chi2_p_value_all_employees": "p-value (хи-квадрат)",
                             "significant_difference": "Различие значимо?"})
    w = write_table(ws, e2, 3, int_cols=["Ошибок в реестре", "Продаж"])
    for r in range(4, w + 1):
        for col in (4, 5, 6):
            ws.cell(r, col).number_format = "0.00%"
        ws.cell(r, 7).number_format = "0.000"

    ws = wb.create_sheet("Модель")
    _title(ws, "ML: предупреждающая модель — предсказывает источник проблемы по операции ETM",
           "Метки получены правилами (слабая разметка), признаки известны на момент проводки, проверка строго по времени, дважды. "
           "Рядом с метриками — тривиальные baseline. Вывод: модель приоритизирует проверку, а не заменяет правила.")
    m = M["metrics"].drop(columns=["fold"]).T.reset_index()
    m.columns = ["Метрика"] + list(M["metrics"]["fold"])
    r_end = write_table(ws, m, 3, widths={"Метрика": 42, m.columns[1]: 24, m.columns[2]: 24}, freeze_col=None)
    r0 = r_end + 3
    ws.cell(r0 - 1, 1, "По классам (тест — последний месяц)").font = Font(name=FONT, bold=True)
    pc = M["per_class"][M["per_class"].test_month == M["metrics"].iloc[0]["test_month"]].drop(columns="test_month")
    r_end = write_table(ws, pc, r0, freeze_col=None)
    r1 = r_end + 3
    ws.cell(r1 - 1, 1, "Важность признаков (падение PR-AUC при перемешивании)").font = Font(name=FONT, bold=True)
    write_table(ws, M["importance"], r1, freeze_col=None)

    # ----------------------------------------------------------------- Кейсы
    ws = wb.create_sheet("Кейсы для защиты")
    _title(ws, "Три реальных случая: сумма → билет → где потерялось → кто проверяет → что исправить")
    r = 3
    for title, agent, what, where, who, fix in build_cases(R):
        ws.cell(r, 1, title).font = Font(name=FONT, bold=True, size=12)
        ws.cell(r, 1).fill = SUB_FILL
        ws.cell(r, 2).fill = SUB_FILL
        for k, vv in [("Субагент", agent), ("Сумма и билет", what), ("Где потерялось", where), ("Кто проверяет", who), ("Что исправить", fix)]:
            r += 1
            ws.cell(r, 1, k).font = Font(name=FONT, bold=True, size=10)
            ws.cell(r, 2, vv).font = Font(name=FONT, size=10)
            ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 32
        r += 2
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 120

    # ----------------------------------------------------------- Вопросы бухгалтеру
    ws = wb.create_sheet("Вопросы бухгалтеру")
    _title(ws, "Конкретные вопросы, которые надо закрыть с бухгалтером")
    write_table(ws, pd.DataFrame(summary["questions"], columns=["№", "Вопрос", "Сколько случаев / KGS", "Как меняется результат, если ответ «да»"]),
                3, widths={"№": 5, "Вопрос": 95, "Сколько случаев / KGS": 30, "Как меняется результат, если ответ «да»": 70},
                wrap_cols=["Вопрос", "Как меняется результат, если ответ «да»"], freeze_col=None)

    # ---------------------------------------------------------------- Допущения
    ws = wb.create_sheet("Допущения и проверки")
    _title(ws, "Допущения, проверки целостности и ограничения")
    items = list(summary["assumptions"]) + [("— проверки целостности (значения из прогона) —", "")] + [(k, str(vv)) for k, vv in R.checks.items()]
    write_table(ws, pd.DataFrame(items, columns=["Пункт", "Значение / пояснение"]), 3,
                widths={"Пункт": 60, "Значение / пояснение": 110}, wrap_cols=["Значение / пояснение"], freeze_col=None)

    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
    wb.save(path)
